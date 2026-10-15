"""SWEET-specific adapter and Excel preservation invariants."""

from pathlib import Path
from types import SimpleNamespace
import importlib.util
import sys

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "dev"))
from update_vehicle_energy_inventories import read_workbook, write_blocks
from stem_vehicle_energy import physical_cng_supply


def test_every_mapped_transport_endpoint_has_a_swiss_workbook_activity():
    import yaml

    pytest.importorskip("premise.vehicle_energy_use")
    from premise.vehicle_energy_use import source_name

    root = Path(__file__).resolve().parents[1]
    config = yaml.safe_load((root / "configuration_file/config.yaml").read_text())
    required = {
        value["ecoinvent alias"]["name"]
        for value in config["production pathways"].values()
        if source_name(value["ecoinvent alias"]["name"])
    }
    present = {
        d["name"]
        for d in read_workbook(root / "inventories/lci-sweet_sure.xlsx")
        if d["location"] == "CH" and d["unit"] == "megajoule"
    }
    assert required and required <= present


def test_swiss_adapter_accepts_duplicate_alias_and_routes_auxiliary_diesel(
    monkeypatch, tmp_path
):
    pytest.importorskip("premise.vehicle_energy_use")
    import premise.external as external
    import premise.vehicle_energy_use as vehicle_energy
    from stem_vehicle_energy import swiss_vehicle_energy

    name = "electricity, used in single deck coach or urban bus"
    dataset = dict(
        name=name,
        location="CH",
        unit="megajoule",
        **{
            "reference product": "electricity, low voltage",
            "exchanges": [
                dict(
                    name="market group for diesel",
                    product="diesel",
                    unit="kilogram",
                    location="NEU",
                    type="technosphere",
                    amount=0.01,
                )
            ],
        },
    )
    config = {
        "production pathways": {
            "coach": {
                "ecoinvent alias": {
                    "name": name + "_FE_coach_electricity",
                    "source name": name,
                }
            }
        }
    }
    monkeypatch.setattr(
        external, "check_inventories", lambda **kwargs: ([dataset], [], config, {})
    )
    monkeypatch.setattr(vehicle_energy, "select_source", lambda *args, **kwargs: None)
    calls = []
    monkeypatch.setattr(
        vehicle_energy, "normalize", lambda ds, *args: calls.append(ds["name"]) or {}
    )
    with swiss_vehicle_energy(tmp_path / "audit.json"):
        external.check_inventories(year=2050)
    assert calls == [name]
    fuel = dataset["exchanges"][0]
    assert (fuel["name"], fuel["location"], fuel["amount"]) == (
        "market for diesel",
        "CH",
        0.01,
    )


def test_xml_replacement_preserves_unrelated_formula_and_cache(tmp_path):
    from zipfile import ZipFile

    path = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "lci"
    rows = [
        ["Activity", "first"],
        ["reference product", "first"],
        ["location", "CH"],
        ["unit", "unit"],
        ["Exchanges"],
        ["name", "amount", "location", "unit", "type", "reference product"],
        ["first", 1, "CH", "unit", "production", "first"],
        [],
        ["Activity", "second"],
        ["reference product", "second"],
        ["location", "CH"],
        ["unit", "unit"],
        ["Exchanges"],
        ["name", "amount", "location", "unit", "type", "reference product"],
        ["second", 1, "CH", "unit", "production", "second"],
        ["retained", 2, "CH", "unit", "technosphere", "retained"],
    ]
    for row in rows:
        ws.append(row)
    ws["B17"] = "=B16*2"
    wb.save(path)
    # Supply an Excel-style cached formula value, which openpyxl would erase.
    with ZipFile(path) as z:
        contents = {n: z.read(n) for n in z.namelist()}
    contents["xl/worksheets/sheet1.xml"] = contents["xl/worksheets/sheet1.xml"].replace(
        b"<f>B16*2</f><v></v>", b"<f>B16*2</f><v>4</v>"
    )
    with ZipFile(path, "w") as z:
        for n, b in contents.items():
            z.writestr(n, b)
    ds = read_workbook(path)[0]
    ds["comment"] = "new provenance expands replaced block"
    write_blocks(path, [ds])
    cached = openpyxl.load_workbook(path, data_only=True).active
    formulas = openpyxl.load_workbook(path, data_only=False).active
    found = [c for row in formulas for c in row if c.data_type == "f"]
    assert len(found) == 1
    assert cached[found[0].coordinate].value == 4
    assert found[0].value == f"=B{found[0].row - 1}*2"


def test_cng_market_is_resolved_to_compression_not_relabelled():
    pytest.importorskip("premise.vehicle_energy_use")
    from premise.vehicle_energy_use import VehicleLookup

    product = "natural gas, medium pressure, vehicle grade"
    physical = dict(
        name="compression",
        **{"reference product": product},
        location="NEU",
        unit="kilogram",
        exchanges=[
            dict(
                name="gas",
                product="natural gas, high pressure",
                location="NEU",
                unit="cubic meter",
                type="technosphere",
                amount=1.36,
            )
        ],
    )
    market = dict(
        name="market",
        **{"reference product": product},
        location="NEU",
        unit="kilogram",
        exchanges=[
            dict(
                name="compression",
                product=product,
                location="NEU",
                unit="kilogram",
                type="technosphere",
                amount=1,
            )
        ],
    )
    assert physical_cng_supply(market, VehicleLookup([market, physical])) is physical
    market["exchanges"][0]["amount"] = 0.9999999944120646
    assert physical_cng_supply(market, VehicleLookup([market, physical])) is physical
    market["exchanges"][0]["amount"] = 0.5
    with pytest.raises(ValueError, match="compression recipe"):
        physical_cng_supply(market, VehicleLookup([market, physical]))
