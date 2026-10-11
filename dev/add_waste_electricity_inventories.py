"""Append electricity-only waste proxies without altering existing formula caches.

The source inventories allocate combustion burdens to useful electricity and
heat. These variants remove that allocation (all burden goes to electricity),
then the normal STEM build applies the separately mapped electrical efficiency.
This is a proxy adaptation, not a newly measured electricity-only plant.
"""

import argparse
from copy import deepcopy
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

VARIANTS = []
for fossil, fuel, fraction in (
    (True, "plastic waste mixture", 0.33),
    (False, "biowaste", 0.29),
):
    for ccs in (False, True):
        source = (
            f"electricity production, at co-generation {'plastic waste' if fossil else 'biowaste'}-fired power plant, post, pipeline 200km, storage 1000m"
            if ccs
            else (
                "electricity production, from incineration of plastic waste mixture, energy allocation"
                if fossil
                else "electricity production, from biowaste, municipal incineration, energy allocation"
            )
        )
        suffix = ", post, pipeline 200km, storage 1000m" if ccs else ""
        VARIANTS.append(
            {
                "source": source,
                "name": f"electricity production, from incineration of {fuel}, electricity only{suffix}",
                "product": f"electricity, from incineration of {fuel}, electricity only"
                + (", with CCS" if ccs else ""),
                "fraction": fraction,
                "pathway": f"waste incineration, {'fossil' if fossil else 'non-fossil'} (electricity only"
                + (", CCS)" if ccs else ")"),
                "variable": f"Electricity generation|Wastes|{'Non Renewable' if fossil else 'Renewable'}|Wastes Incineration (electric only)"
                + (" CCS" if ccs else ""),
                "efficiency": "Efficiency of power plants (electric)|Wastes|Wastes Incineration (electric only)"
                + (" CCS" if ccs else ""),
            }
        )


def blocks(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    assert workbook.sheetnames[0] == "lci"
    rows = [list(r) for r in workbook["lci"].values]
    workbook.close()
    starts = [i for i, r in enumerate(rows) if r[0] == "Activity"] + [len(rows)]
    return {rows[a][1]: rows[a:b] for a, b in zip(starts, starts[1:])}


def variant_rows(source, spec):
    rows = deepcopy(source)
    note = (
        f"Electricity-only proxy derived from '{spec['source']}'. All energy-production "
        f"burdens assigned to electricity: divide non-production exchanges by the "
        f"source electricity allocation factor {spec['fraction']}. No useful heat "
        "coproduct. Retain the source's declared physical electrical efficiency "
        "before applying the matching absolute STEM efficiency. CCS capture and "
        "storage, residual emissions and biomass uptake retain their source ratios."
    )
    exchanges = False
    for row in rows:
        if row[0] == "Exchanges":
            exchanges = True
        elif not exchanges:
            if row[0] == "Activity":
                row[1] = spec["name"]
            elif row[0] == "reference product":
                row[1] = spec["product"]
            elif row[0] == "comment":
                row[1] = note + " Source documentation: " + str(row[1] or "")
        elif row[5] == "production":
            row[0], row[7] = spec["name"], spec["product"]
            assert row[1] == 1
        elif row[5] in ("biosphere", "technosphere"):
            assert isinstance(row[1], (float, int)), row
            row[1] /= spec["fraction"]
            row[6] = (str(row[6]) + " " if row[6] else "") + (
                f"Electricity-only proxy: source amount / {spec['fraction']}; "
                "100% burden to electricity, before STEM efficiency adjustment."
            )
            # Amounts are explicit: no replacement metadata may overwrite them.
            assert not any(row[8:]), row
    return rows


def append_variants(path):
    original = blocks(path)
    present = [v["name"] in original for v in VARIANTS]
    if all(present):
        return 0
    if any(present):
        raise ValueError("Only some electricity-only variants already exist")
    additions = []
    for spec in VARIANTS:
        additions.extend(variant_rows(original[spec["source"]], spec))
        additions.append([None] * 15)

    # openpyxl is used to read/validate; saving the entire workbook would clear
    # cached formula values needed by bw2io. Append literal rows to sheet XML.
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", ns)
    q = lambda tag: f"{{{ns}}}{tag}"
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        data = root.find(q("sheetData"))
        last = max(int(r.attrib["r"]) for r in data) + 1
        for index, values in enumerate(additions, last + 1):
            row = ET.SubElement(data, q("row"), r=str(index))
            for column, value in enumerate(values, 1):
                if value is None:
                    continue
                cell = ET.SubElement(
                    row, q("c"), r=f"{get_column_letter(column)}{index}"
                )
                if isinstance(value, (float, int)):
                    ET.SubElement(cell, q("v")).text = repr(value)
                else:
                    cell.set("t", "inlineStr")
                    ET.SubElement(ET.SubElement(cell, q("is")), q("t")).text = str(
                        value
                    )
        root.find(q("dimension")).set("ref", f"A1:O{index}")
        with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as temp:
            temporary = Path(temp.name)
        with ZipFile(temporary, "w") as output:
            for item in archive.infolist():
                output.writestr(
                    item,
                    (
                        ET.tostring(root)
                        if item.filename == "xl/worksheets/sheet1.xml"
                        else archive.read(item.filename)
                    ),
                )
    temporary.replace(path)
    return len(VARIANTS)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    print(
        f"Added {append_variants(parser.parse_args().workbook)} electricity-only inventories"
    )
